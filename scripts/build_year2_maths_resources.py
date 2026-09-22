"""Year 2: five illustrated questions on one A4 pupil page, separate answers."""
import json,math,re,hashlib,shutil,argparse
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import RGBColor
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor,white
from pypdf import PdfReader
from build_year1_picture_worksheets import text,para,box,line,shape,objects,frame,numberline,tens,solid,clockface,coin,answer_text,BLUE,GREEN,PURPLE,ORANGE,INK,LINE
ROOT=Path(__file__).resolve().parents[1];QA=ROOT/'.qa/year2-maths';W,H=A4
REF=ROOT/'lessons/year-1-maths/week-5/add-by-counting-on/editable-teacher-plan.docx'

def arrow(c,x,y,s,d,color=BLUE):
    dx,dy=[(0,1),(1,0),(0,-1),(-1,0)][d]
    line(c,x-dx*s/2,y-dy*s/2,x+dx*s/2,y+dy*s/2,color,2)
    for side in [-1,1]:line(c,x+dx*s/2,y+dy*s/2,x+dx*(s/2-7)+dy*5*side,y+dy*(s/2-7)-dx*5*side,color,2)

def flat(c,n,x,y,s=40,color=BLUE):
    if n in ['pentagon','hexagon']:
        count=5 if n=='pentagon' else 6;p=c.beginPath()
        for i in range(count):
            a=math.pi/2+i*math.pi*2/count;pt=(x+s/2+s/2*math.cos(a),y+s/2+s/2*math.sin(a))
            (p.moveTo if i==0 else p.lineTo)(*pt)
        p.close();c.setFillColor(HexColor(color));c.setStrokeColor(HexColor(INK));c.drawPath(p,fill=1,stroke=1)
    else:shape(c,n,x,y,s,color)

def solid2(c,n,x,y,s=42):
    if n=='triangular prism':
        pts=[(x,y),(x+s,y),(x+s*.5,y+s)];p=c.beginPath();p.moveTo(*pts[0]);p.lineTo(*pts[1]);p.lineTo(*pts[2]);p.close();c.setFillColor(HexColor('#BFDBFE'));c.setStrokeColor(HexColor(BLUE));c.drawPath(p,fill=1,stroke=1)
        for a,b in [(pts[1],(x+s+18,y+13)),(pts[2],(x+s*.5+18,y+s+13)),((x+s+18,y+13),(x+s*.5+18,y+s+13))]:line(c,*a,*b,BLUE)
    else:solid(c,n.replace('square based','square-based'),x,y,s)

def base(c,n,x,y,w=100,h=70):
    t,o=divmod(n,10);unit=min(5,(w-45)/max(t,1)/1.6)
    for j in range(t):
        for k in range(10):
            c.setFillColor(HexColor(BLUE));c.setStrokeColor(white);c.rect(x+j*(unit+2),y+k*5,unit,5,fill=1,stroke=1)
    for j in range(o):shape(c,'circle',x+t*(unit+2)+8+j%3*9,y+j//3*10,6,GREEN)

def counters(c,n,x,y,w,h,color=BLUE):
    cols=min(10,max(1,n),max(1,int(w/10)));rows=max(1,math.ceil(n/cols));s=min(15,(w-4)/cols-3,(h-4)/rows-2)
    assert s>=5,(n,w,h,s)
    for i in range(n):shape(c,'circle',x+i%cols*(s+3),y+h-s-i//cols*(s+2),s,color)

def model(c,m,x,y,w=215,h=84,solved=False,level='expected'):
    k,mode=m['kind'],m['mode']
    if k=='place':
        n=m['n']
        if mode in ['line','between','compareMissing']:
            lo=m['lo'];hi=m['hi'];numberline(c,x+8,y+30,w-16,lo,hi)
        elif mode=='order':
            for i,n in enumerate(m['values']):box(c,x+i*w/4,y+25,w/4-5,32,'#EFF6FF');text(c,n,x+(i+.5)*w/4-2,y+35,15,center=True)
        elif mode in ['compare','symbols']:
            base(c,n,x,y+20,w/2-8);base(c,m['b'],x+w/2+5,y+20,w/2-8)
        else:base(c,n,x+10,y+15,w*.7);text(c,'tens  ones',x+100,y+65,10,BLUE);text(c,'___   ___',x+103,y+30,12)
    elif k=='sequence':
        for i,n in enumerate(m['values']):
            box(c,x+i*w/5,y+26,w/5-4,32,'#EFF6FF');text(c,n if solved or i!=2 else '?',x+(i+.5)*w/5-2,y+36,13,center=True)
    elif k=='bonds':
        total=m['total'];part=m['part'];text(c,'whole '+str(total),x+w/2,y+69,12,BLUE,True,True)
        box(c,x+5,y+22,w*.46,30,'#EFF6FF');text(c,part,x+w*.25,y+31,14,center=True);box(c,x+w*.53,y+22,w*.44,30);text(c,total-part if solved else '?',x+w*.75,y+31,14,center=True)
    elif k=='calc':
        a,b=m['a'],m['b'];sub=m['op']=='−'
        if mode in ['missing','family']:
            whole=a if sub else m['total'];known=m['total'] if sub else a
            text(c,'whole '+str(whole),x+w/2,y+66,12,BLUE,True,True);box(c,x+8,y+18,w-16,34);line(c,x+w/2,y+18,x+w/2,y+52,BLUE)
            text(c,known,x+w/4,y+29,15,center=True);text(c,b if solved or mode=='family' else '?',x+w*.75,y+29,15,center=True)
        elif mode in ['three','twoStep']:
            vals=[a,b,m['c']]
            for i,n in enumerate(vals):box(c,x+i*w/3,y+28,w/3-6,30,'#EFF6FF');text(c,n,x+(i+.5)*w/3-3,y+38,15,center=True)
            text(c,'+ then '+('−' if mode=='twoStep' else '+'),x+w/2,y+8,12,center=True)
        elif a<20 and b<20:
            objects(c,a,x+5,y+12,w*.45,65,BLUE);text(c,m['op'],x+w/2,y+38,16,center=True);objects(c,b,x+w*.58,y+12,w*.4,65,GREEN)
        else:
            base(c,a,x+5,y+15,w*.45);text(c,m['op'],x+w/2,y+40,18,center=True);base(c,b,x+w*.57,y+15,w*.4)
    elif k in ['groups','division']:
        total=m['total'];groups=m['groups'];size=m['size']
        if mode=='parity':
            for j in range(total):shape(c,'circle',x+8+(j//2)*17,y+30+(j%2)*15,10,BLUE)
        elif k=='division':
            counters(c,total,x+4,y+30,w-8,62,BLUE)
            g=groups if m.get('share') else 3
            for j in range(g):box(c,x+j*w/g,y+3,w/g-5,26)
            if not m.get('share'):text(c,'Draw more groups if needed.',x,y-9,9,BLUE)
        else:
            sz=min(12,(w-12)/max(size,1)-4,70/max(groups,1)-2)
            for r in range(groups):
                for col in range(size):shape(c,'circle',x+8+col*(sz+4),y+7+r*(sz+2),sz,[BLUE,GREEN][r%2])
    elif k=='money':
        if m.get('notes'):
            for j,v in enumerate(m['notes']):box(c,x+j*92,y+25,82,42,'#E1F1EB');text(c,'£'+str(v),x+j*92+41,y+39,17,center=True)
        elif mode=='change':coin(c,100,x+10,y+23,38);box(c,x+90,y+27,85,29,'#FFF7ED');text(c,str(m['price'])+'p',x+132,y+36,14,center=True)
        elif mode=='twoItems':
            for j,v in enumerate([m['price'],m['second']]):box(c,x+j*85,y+32,74,29,'#FFF7ED');text(c,str(v)+'p',x+j*85+37,y+41,13,center=True)
            text(c,'Pay £1',x+65,y+9,12,center=True)
        else:
            for j,v in enumerate(m['coins']):coin(c,v,x+4+j*44,y+30,35)
            text(c,'Coin value models',x+5,y+9,8)
    elif k=='measure':
        a=m.get('a',0)
        if mode in ['cm','offset','lengthCompare']:
            start=m.get('start',0);end=m.get('end',a);hi=max(10,end+1);unit=(w-18)/hi
            numberline(c,x+7,y+24,w-18,0,hi)
            c.setFillColor(HexColor('#93C5FD'));c.rect(x+7+start*unit,y+32,(end-start)*unit,12,fill=1,stroke=0)
            if mode=='lengthCompare':c.setFillColor(HexColor('#A7D9C8'));c.rect(x+7,y+50,m['b']*unit,12,fill=1,stroke=0);text(c,'B',x+w-5,y+51,9)
            text(c,'cm',x+w-17,y+67,9)
        elif mode in ['mass','capacity','temperature']:
            if mode=='temperature':
                xx=x+55;bottom=y+6;scale=2
                c.setStrokeColor(HexColor(BLUE));c.rect(xx,bottom,13,60,fill=0);c.setFillColor(HexColor('#ED8A76'));c.rect(xx+2,bottom,9,m['value']*scale,fill=1,stroke=0)
                for n in range(0,31,2):line(c,xx+13,bottom+n*scale,xx+20,bottom+n*scale,INK,.6)
                for n in range(0,31,10):text(c,n,xx+24,bottom+n*scale-3,8)
                text(c,'°C',xx+62,bottom+27,12)
            else:
                mx=m['max'];numberline(c,x+8,y+28,w-16,0,mx,m['step']);at=x+8+m['value']/mx*(w-16);arrow(c,at,y+54,26,2,ORANGE);text(c,m['unit'],x+w-18,y+67,10)
        elif mode=='metres':
            for j in range(a):box(c,x+5+j*56,y+36,51,24,'#DBEAFE');text(c,'1 m',x+30+j*56,y+43,10,center=True)
            text(c,'and '+str(m['b'])+' cm',x+10,y+13,11)
        elif mode in ['lengthProblem','massCompare']:
            for j,v in enumerate([m['a'],m['b']]):box(c,x+7,y+43-j*34,w-20,25,'#EFF6FF' if j==0 else '#F0FDF4');text(c,str(v)+' '+m['unit'],x+w/2,y+50-j*34,12,center=True)
        else:
            for j,unit in enumerate(m['choices']):box(c,x+15+j*85,y+25,74,36,'#EFF6FF');text(c,unit,x+52+j*85,y+37,16,center=True)
    elif k=='fraction':
        den=m['den'];num=m['num'];total=m['total']
        if mode in ['equal','equivalent','count'] or m.get('shape'):
            if mode=='count':
                for i in range(5):text(c,'?' if i==2 and not solved else '0' if i==0 else f'{i}/{den}',x+15+i*40,y+36,14,center=True)
            else:
                for row,d in enumerate([2,4] if mode=='equivalent' else [den]):
                    xx=x+7;unit=(w-14)/d
                    for j in range(d):
                        ww=unit*(.5 if j==0 else 1.5 if j==d-1 else 1) if mode=='equal' and not m['equal'] else unit
                        c.setFillColor(white);c.setStrokeColor(HexColor(PURPLE));c.rect(xx,y+13+row*32,ww,26,fill=1,stroke=1);xx+=ww
        else:
            if mode=='whole':
                box(c,x+7,y+23,w-14,39)
                for j in range(1,den):line(c,x+7+j*(w-14)/den,y+23,x+7+j*(w-14)/den,y+62,PURPLE)
                counters(c,total//den,x+12,y+29,(w-14)/den-9,26,PURPLE)
            else:
                objects(c,total,x+5,y+33,w-10,60,PURPLE)
                for j in range(den):box(c,x+5+j*(w-8)/den,y+3,(w-14)/den-4,25)
    elif k=='shape':
        if mode in ['solid','faces']:solid2(c,m['name'],x+50,y+15,48)
        elif mode=='sort':
            for j,n in enumerate(m['names']):flat(c,n,x+4+j*42,y+27,24,[BLUE,GREEN,PURPLE][j%3])
        else:flat(c,m['name'],x+55,y+16,52)
    elif k=='direction':
        if mode in ['position','route']:
            for row in range(5):
                for col in range(5):c.setStrokeColor(HexColor(LINE));c.rect(x+45+col*14,y+row*14,14,14,fill=0)
            shape(c,'circle',x+49,y+4,6,BLUE)
        elif mode=='pattern':
            for j,d in enumerate(m['values']):arrow(c,x+12+j*32,y+42,25,d)
        else:arrow(c,x+42,y+40,45,m['start']);box(c,x+106,y+12,70,58);text(c,'finish',x+141,y+2,9,center=True)
    elif k=='time':
        if mode=='duration':
            for j in range(m['hours']):clockface(c,x+35+j*65,y+43,28,1,0)
        else:clockface(c,x+70,y+42,39,m['hour'] if mode!='draw' or solved else None,m['minute'])
    elif k=='data':
        vals=m['values'];key=m['key'];labels=m['labels']
        for j,label in enumerate(labels):
            yy=y+64-j*24;text(c,label,x,yy,9)
            if mode=='table':text(c,vals[j],x+95,yy,13)
            elif mode=='tally':
                for k2 in range(vals[j]):
                    xx=x+48+(k2//5)*27+(k2%5)*5
                    if k2%5==4:line(c,xx-22,yy-2,xx,yy+10,INK,1)
                    else:line(c,xx,yy-3,xx,yy+11,INK,1)
            elif j==2 and mode in ['draw','complete'] and not solved:line(c,x+50,yy,x+w-4,yy)
            else:
                for k2 in range(vals[j]//key):shape(c,'square' if mode in ['block','draw','total'] else 'circle',x+48+k2*17,yy-3,12,[BLUE,GREEN,PURPLE][j])
        if mode not in ['tally','table']:text(c,'Each '+('block' if mode in ['block','draw','total'] else 'circle')+f' = {key}',x+50,y-8,8)
    else:raise ValueError(k)

def header(c,l,label):
    text(c,'BrightPath  |  Year 2 Maths',32,H-30,10,GREEN,True)
    para(c,l['title'],32,H-42,W-64,18,46)
    text(c,label,32,H-101,12,GREEN,True)
    line(c,30,28,W-30,28);text(c,'BrightPath',32,15,8)

def worksheet(l,level,folder):
    p=folder/(level+'-worksheet.pdf');c=canvas.Canvas(str(p),pagesize=A4,invariant=1,pageCompression=1);c.setTitle(l['title']+' '+level+' worksheet');c.setAuthor('BrightPath')
    header(c,l,{'lower':'Supported practice','expected':'Your turn','higher':'Think and explain'}[level]);text(c,'Name: ______________________',320,H-101,10)
    qs=l[level]
    for i,q in enumerate(qs):
        top=H-124-i*136;bottom=top-124;line(c,30,bottom,W-30,bottom)
        text(c,str(i+1),32,top-15,13,GREEN,True)
        para(c,q['q'],52,top,W*.46,12,78)
        model(c,q['model'],330,bottom+24,228,85,False,level)
        line(c,52,bottom+15,302,bottom+15)
    c.showPage();header(c,l,'Answers - print only if needed')
    for i,q in enumerate(qs):
        top=H-128-i*130;text(c,str(i+1),33,top-14,13,GREEN,True);para(c,answer_text(q['a']),55,top,W-92,13,105)
        line(c,33,top-115,W-33,top-115)
    c.showPage();c.save();assert len(PdfReader(p).pages)==2

def preteach(l,folder):
    p=folder/'pre-teach.pdf';c=canvas.Canvas(str(p),pagesize=A4,invariant=1,pageCompression=1);header(c,l,'Pre-teach: a short supported practice')
    para(c,l['preteach']['focus'],32,H-127,W-64,13,48)
    for i,q in enumerate(l['preteach']['questions']):
        top=H-190-i*170;text(c,str(i+1),32,top-12,13,GREEN,True);para(c,q['q'],54,top,255,13,75);model(c,q['model'],333,top-105,220,90,False,'lower');line(c,54,top-125,300,top-125)
    para(c,'Adult: let the pupil move objects first. Ask them to describe the model before recording.',32,95,W-64,11,38);c.showPage();header(c,l,'Pre-teach answer guide')
    for i,q in enumerate(l['preteach']['questions']):para(c,f"{i+1}. {q['a']}",35,H-140-i*110,W-70,14,92)
    c.showPage();c.save()

def setp(p,value):
    # Preserve the reference paragraph and its styles, replacing only the text.
    if p.runs:
        p.runs[0].text=value
        for r in p.runs[1:]:r.text=''
    else:p.add_run(value)

def plan(l,folder):
    doc=Document(REF);ps=doc.paragraphs
    setp(ps[1],l['title']);setp(ps[2],f"Year 2 Maths | {l['term']} | Week {l['week']} | Day {l['day']} | Editable teacher plan")
    for p,value in zip(ps[4:7],l['successCriteria']):setp(p,value)
    setp(ps[14],l['misconception'])
    info=doc.tables[0]
    for row,value in zip(info.rows[2:],[l['objective'],l['prior'],', '.join(l['vocabulary']),l['resources'],l['title']]):row.cells[1].text=value
    seq=doc.tables[3]
    values=[l['warmup'],'\n'.join(f'{i+1}. {s}' for i,s in enumerate(l['teacherModel'])),l['guided'],l['independent'],l['plenary']]
    for row,value in zip(seq.rows[1:],values):row.cells[2].text=value
    # The source answer guide has fixed heading slots and 3/5/5/5 answers.
    groups={'Pre-teach':l['preteach']['questions'],'Lower support':l['lower'],'Expected':l['expected'],'Higher challenge':l['higher']}
    active=None;idx=0
    for p in doc.paragraphs:
        if p.text in groups:active=groups[p.text];idx=0
        elif active is not None and re.match(r'^\d+\.',p.text):
            setp(p,f"{idx+1}. {active[idx]['a']}");idx+=1
            if idx==len(active):active=None
    for el in doc.element.iter():
        for name,value in list(el.attrib.items()):
            if value.upper() in ['B91C1C','DC2626']:el.set(name,'167D8D')
            elif value.upper()=='FEF2F2':el.set(name,'F0F8FA')
    doc.save(folder/'editable-teacher-plan.docx')

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--only',nargs='*');args=ap.parse_args()
    rows=json.loads((QA/'lessons.json').read_text(encoding='utf-8'));report=[]
    for l in rows:
        id=f"2-{l['week']}-{l['day']}"
        if args.only and id not in args.only:continue
        folder=QA/'resources'/id;folder.mkdir(parents=True,exist_ok=True)
        for level in ['lower','expected','higher']:worksheet(l,level,folder)
        preteach(l,folder);plan(l,folder)
        report.append({'id':id,'worksheets':3,'questionsPerPupilPage':5,'pupilPagesPerWorksheet':1})
        print('BUILT '+id,flush=True)
    (QA/'resource-build-report.json').write_text(json.dumps(report,indent=2))

if __name__=='__main__':main()
