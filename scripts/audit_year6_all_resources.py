import json
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / "public" / "lessons" / "year-6-maths"


def check(condition, message):
    if not condition:
        raise RuntimeError(message)


def inspect_pptx(path):
    with zipfile.ZipFile(path) as archive:
        slides = sorted(name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
        check(len(slides) == 12, f"{path} has {len(slides)} slides")
        click_effects = 0
        for name in slides:
            data = archive.read(name)
            check(b"[[CLICK]]" not in data, f"animation marker remains in {path} {name}")
            click_effects += data.count(b'nodeType="clickEffect"')
        slide4 = archive.read("ppt/slides/slide4.xml")
        slide4_clicks = slide4.count(b'nodeType="clickEffect"')
        check(click_effects >= 25, f"too few click animations in {path}: {click_effects}")
        detailed_division = b"Useful multiples" in slide4
        check(slide4_clicks >= (7 if detailed_division else 5), f"worked model is not sufficiently staged in {path}: {slide4_clicks}")


def inspect_docx(path, expected_tables, required_text):
    with zipfile.ZipFile(path) as archive:
        check("word/document.xml" in archive.namelist(), f"invalid Word package: {path}")
    document = Document(path)
    check(len(document.tables) == expected_tables, f"{path} has {len(document.tables)} tables")
    content = " ".join(paragraph.text for paragraph in document.paragraphs)
    content += " " + " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    for value in required_text:
        check(value.lower() in content.lower(), f"{value!r} missing from {path}")


def inspect_term(term):
    root = PUBLIC / term
    lessons = json.loads((root / f"year6-{term}-lessons.json").read_text(encoding="utf-8"))
    check(len(lessons) == 50, f"{term} has {len(lessons)} lesson records")
    for item in lessons:
        folder = root / f"week-{item['week']}" / item["slug"]
        powerpoint = folder / "teaching-powerpoint-v3.pptx"
        if not powerpoint.exists():
            powerpoint = folder / "teaching-powerpoint-v1.pptx"
        inspect_pptx(powerpoint)
        plan = folder / "editable-teacher-plan.docx"
        inspect_docx(plan, 4, ["resources", "SEND", "ADHD", "ODD", "assessment"])
        check(len(PdfReader(str(folder / "pre-teach.pdf")).pages) == 2, f"pre-teach page count failed: {folder}")
        check(len(PdfReader(str(folder / "differentiated-worksheets.pdf")).pages) == 6, f"worksheet page count failed: {folder}")
        check(len(list((folder / "preview" / "powerpoint").glob("slide-*.png"))) == 12, f"PowerPoint previews failed: {folder}")
        check(len(list((folder / "preview" / "preteach").glob("page-*.png"))) == 2, f"pre-teach previews failed: {folder}")
        check(len(list((folder / "preview" / "worksheets").glob("page-*.png"))) == 6, f"worksheet previews failed: {folder}")
    print(f"{term}: 50 lessons verified")


def inspect_evaluation(term):
    folder = PUBLIC / "evaluations" / term
    metadata = json.loads((folder / "evaluation.json").read_text(encoding="utf-8"))
    check(metadata["totalMarks"] == 40, f"{term} evaluation total is not 40")
    check([(b["minimum"], b["maximum"]) for b in metadata["bands"]] == [(0, 15), (16, 23), (24, 31), (32, 40)], f"{term} score bands differ")
    check(len(PdfReader(str(folder / "term-assessment-and-mark-scheme.pdf")).pages) == 5, f"{term} evaluation PDF page count")
    inspect_docx(folder / "editable-teacher-evaluation-record.docx", 3, ["Score /40", "Cuspy", "Intervention", "Review evidence", "[30]"])
    check(len(list((folder / "preview" / "assessment").glob("page-*.png"))) == 5, f"{term} assessment previews")
    check(len(list((folder / "preview" / "teacher-record").glob("page-*.png"))) == 2, f"{term} record previews")
    print(f"{term} evaluation: assessment, bands and editable record verified")


if __name__ == "__main__":
    for selected in ("autumn", "spring", "summer"):
        inspect_term(selected)
        inspect_evaluation(selected)
    print("Year 6 Maths resource audit passed: 150 lessons and 3 term evaluations")
