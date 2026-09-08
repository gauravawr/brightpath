import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
WEEK = int(sys.argv[1]) if len(sys.argv) > 1 else 1
DATA = ROOT / f"public/lessons/year-1-maths/week-{WEEK}/week{WEEK}-lessons.json"
OUT = DATA.parent
RED = "B91C1C"
RED_MID = "DC2626"
PALE = "FEF2F2"
SLATE = "1E293B"
MUTED = "64748B"
LINE = "E2E8F0"
pdfmetrics.registerFont(TTFont("ArialBP", r"C:\Windows\Fonts\arial.ttf"))
pdfmetrics.registerFont(TTFont("ArialBPBold", r"C:\Windows\Fonts\arialbd.ttf"))
pdfmetrics.registerFont(TTFont("SegoeSymbolBP", r"C:\Windows\Fonts\seguisym.ttf"))


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=SLATE, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_doc_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(str(item))


def add_info_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(4.4)
    table.columns[1].width = Cm(12.5)
    for idx, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        shade(cells[0], RED if idx % 2 == 0 else RED_MID)
        set_cell_text(cells[0], label, True, "FFFFFF")
        shade(cells[1], "FFFFFF" if idx % 2 == 0 else PALE)
        set_cell_text(cells[1], value)
    return table


def add_adaptation_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Pupil initials / group", "Barrier or need", "Adaptation, reasonable adjustment or adult support"]
    for cell, title in zip(table.rows[0].cells, headers):
        shade(cell, RED)
        set_cell_text(cell, title, True, "FFFFFF")
    set_repeat_table_header(table.rows[0])
    for _ in range(5):
        cells = table.add_row().cells
        for cell in cells:
            set_cell_text(cell, "\n\n")
    return table


def create_plan(lesson):
    folder = OUT / lesson["slug"]
    folder.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Cm(1.45)
    sec.left_margin = sec.right_margin = Cm(1.6)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("BRIGHTPATH PRIMARY LEARNING")
    r.bold = True; r.font.name = "Arial"; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(RED)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run(lesson["title"])
    r.bold = True; r.font.name = "Arial"; r.font.size = Pt(23); r.font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph(f"Year 1 Maths | Autumn | Week {WEEK} | Day {lesson['day']} | Editable teacher plan")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    add_info_table(doc, [
        ("Teacher / class", ""), ("Date / time", ""), ("Learning objective", lesson["objective"]),
        ("Prior learning", lesson["prior"]), ("Vocabulary", ", ".join(lesson["vocabulary"])),
        ("Resources", lesson["resources"]), ("National Curriculum link", "Count, read, write and use numbers; use concrete objects and pictorial representations; explain mathematical thinking."),
    ])
    add_doc_heading(doc, "Success criteria", 2)
    add_bullets(doc, lesson["successCriteria"])

    doc.add_page_break()
    add_doc_heading(doc, "Editable inclusion and adaptation plan", 1)
    p = doc.add_paragraph("Complete before teaching. Record initials rather than full pupil names. Adapt the pitch, representation, language, pace and outcome to remove barriers while keeping the mathematical goal ambitious.")
    p.runs[0].italic = True
    add_adaptation_table(doc)
    add_doc_heading(doc, "Whole-class adjustments", 2)
    add_info_table(doc, [
        ("SEND / cognition and learning", "Use concrete objects, one instruction at a time, visual examples, reduced item count and overlearning."),
        ("Communication / EAL", "Pre-teach vocabulary with objects and gestures; model a full sentence stem; allow rehearsal with a partner."),
        ("Sensory / physical", "Offer larger objects, uncluttered workspace, alternative recording and movement breaks as needed."),
        ("Higher prior attainment", "Ask for a second method, proof, counterexample or generalisation; require precise mathematical language."),
        ("Teacher edits", ""), ("Adult deployment", ""),
    ])

    doc.add_page_break()
    add_doc_heading(doc, "Teaching sequence (60 minutes)", 1)
    sequence = [
        ("0-5 min", "Revisit and assess", lesson["warmup"], "Listen for secure prerequisite language."),
        ("5-20 min", "Explicit modelling", "\n".join(f"{i+1}. {step}" for i, step in enumerate(lesson["teacherModel"])), "Think aloud; vary the example and non-example."),
        ("20-35 min", "Guided practice", lesson["guided"], "Question pupils before accepting choral answers."),
        ("35-55 min", "Independent practice", lesson["independent"], "Choose the worksheet by current need, not a fixed label."),
        ("55-60 min", "Plenary / exit check", lesson["plenary"], "Record who is ready, needs pre-teach or needs challenge."),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Time", "Phase", "Teacher and pupil activity", "Assessment / adaptation"]):
        shade(cell, RED); set_cell_text(cell, text, True, "FFFFFF")
    set_repeat_table_header(table.rows[0])
    for idx, row in enumerate(sequence):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            shade(cell, "FFFFFF" if idx % 2 == 0 else PALE); set_cell_text(cell, value, idx == 0)
    add_doc_heading(doc, "Key misconception and response", 2)
    doc.add_paragraph(lesson["misconception"])
    add_doc_heading(doc, "Assessment notes / next lesson", 2)
    doc.add_paragraph("\n\n\n")

    doc.add_page_break()
    add_doc_heading(doc, "Pre-teach and worksheet answer guide", 1)
    for label, questions in (("Pre-teach", lesson["preteach"]["questions"]), ("Lower support", lesson["lower"]), ("Expected", lesson["expected"]), ("Higher challenge", lesson["higher"])):
        add_doc_heading(doc, label, 2)
        for idx, item in enumerate(questions, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.add_run(f"{idx}. ").bold = True
            p.add_run(item["a"])
    add_doc_heading(doc, "Reflection after teaching", 2)
    add_info_table(doc, [("What pupils understood", ""), ("Who needs further support", ""), ("Change before teaching again", "")])

    path = folder / "editable-teacher-plan.docx"
    doc.save(path)
    return path


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("TitleBP", parent=styles["Title"], fontName="ArialBPBold", fontSize=20, leading=23, textColor=HexColor("#0F172A"), alignment=TA_LEFT, spaceAfter=5),
        "meta": ParagraphStyle("MetaBP", parent=styles["Normal"], fontName="ArialBPBold", fontSize=9, textColor=HexColor("#B91C1C"), spaceAfter=6),
        "body": ParagraphStyle("BodyBP", parent=styles["BodyText"], fontName="ArialBP", fontSize=10, leading=14, textColor=HexColor("#334155")),
        "question": ParagraphStyle("QuestionBP", parent=styles["BodyText"], fontName="SegoeSymbolBP", fontSize=10.5, leading=14, textColor=HexColor("#1E293B"), spaceAfter=2),
        "small": ParagraphStyle("SmallBP", parent=styles["BodyText"], fontName="ArialBP", fontSize=8.5, leading=11, textColor=HexColor("#64748B")),
    }


def header_footer(canvas, doc, label):
    canvas.saveState()
    canvas.setFillColor(HexColor("#B91C1C")); canvas.rect(0, A4[1]-9*mm, A4[0], 9*mm, fill=1, stroke=0)
    canvas.setFillColor(HexColor("#FFFFFF")); canvas.setFont("ArialBPBold", 9); canvas.drawString(15*mm, A4[1]-6*mm, "BRIGHTPATH PRIMARY LEARNING")
    canvas.setFillColor(HexColor("#64748B")); canvas.setFont("ArialBP", 8); canvas.drawString(15*mm, 9*mm, label); canvas.drawRightString(A4[0]-15*mm, 9*mm, "Free classroom resource")
    canvas.restoreState()


def create_pupil_pdf(lesson, kind, questions, subtitle):
    folder = OUT / lesson["slug"]
    path = folder / f"{kind}.pdf"
    st = pdf_styles()
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=15*mm, rightMargin=15*mm, topMargin=17*mm, bottomMargin=16*mm)
    story = [
        Spacer(1, 2*mm), Paragraph(lesson["title"], st["title"]),
        Paragraph(f"YEAR 1 MATHS  |  AUTUMN WEEK {WEEK}, DAY {lesson['day']}  |  {subtitle.upper()}", st["meta"]),
        Table([[Paragraph("Name: ____________________________________", st["body"]), Paragraph("Date: __________________", st["body"])]], colWidths=[120*mm, 55*mm], style=TableStyle([("BOX",(0,0),(-1,-1),.7,HexColor("#E2E8F0")),("BACKGROUND",(0,0),(-1,-1),HexColor("#FEF2F2")),("PADDING",(0,0),(-1,-1),7)])),
        Spacer(1, 5*mm),
    ]
    for idx, item in enumerate(questions, 1):
        q = item["q"] if isinstance(item, dict) else item
        box = Table([[Paragraph(f"<b>{idx}.</b> {q}", st["question"])], [""]], colWidths=[180*mm], rowHeights=[None, 19*mm], style=TableStyle([("BOX",(0,0),(-1,-1),.8,HexColor("#CBD5E1")),("LINEBELOW",(0,0),(-1,0),.5,HexColor("#E2E8F0")),("BACKGROUND",(0,0),(-1,0),HexColor("#FFFFFF")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),7)]))
        story += [KeepTogether(box), Spacer(1, 3*mm)]
    story.append(Paragraph("Check: I used objects, pictures or words to show my thinking.", st["small"]))
    doc.build(story, onFirstPage=lambda c,d: header_footer(c,d,subtitle), onLaterPages=lambda c,d: header_footer(c,d,subtitle))
    return path


def create_preteach_pdf(lesson):
    items = [{"q": q["q"]} for q in lesson["preteach"]["questions"]]
    items.insert(0, {"q": "Say and point: " + ", ".join(lesson["vocabulary"]) + "."})
    return create_pupil_pdf(lesson, "pre-teach", items[:5], "Pre-teach: small-group preparation")


def main():
    lessons = json.loads(DATA.read_text(encoding="utf-8"))
    made = []
    for lesson in lessons:
        made.append(create_plan(lesson))
        made.append(create_preteach_pdf(lesson))
        made.append(create_pupil_pdf(lesson, "lower-worksheet", lesson["lower"][:5], "Lower support worksheet"))
        made.append(create_pupil_pdf(lesson, "expected-worksheet", lesson["expected"][:5], "Expected worksheet"))
        made.append(create_pupil_pdf(lesson, "higher-worksheet", lesson["higher"][:5], "Higher challenge worksheet"))
    print(f"Created {len(made)} resources")
    for path in made:
        print(path.relative_to(ROOT))


if __name__ == "__main__":
    main()
